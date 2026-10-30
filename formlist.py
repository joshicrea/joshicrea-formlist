"""Build a list of companies with contact forms for an area and industries.

Runtime: Windows (target) / cross-platform, Python 3.9+.
Usage:
  python formlist.py                                  (interactive, Japanese prompts)
  python formlist.py --area 東京都世田谷区 --industry ホテル・旅館 --industry 百貨店 --limit 200
  python formlist.py --input 会社リスト.csv           (URLs or company names, one per line)
Outputs (in 出力/): フォームリスト_<area>_<time>.xlsx and .csv
Columns A..V are identical to 05 企業リスト台帳, then added check columns.
"""
import argparse
import csv
import datetime as dt
import hashlib
import re
import unicodedata
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Lock
from urllib.parse import quote_plus, urlparse

import openpyxl
from openpyxl.styles import Font, PatternFill

import google_places_text_source
import hp_finder
import hp_verify
from limited_http_budget import HttpBudgetExceeded
import size_guess
import subareas
import map_browser_source
import ng_list
from ng_state_store import NgStateStore
import industries
import osm_source
import overture_source
import portal_source
import site_checker
import state

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "出力"

LEDGER_COLS = ['No', '企業名', '業種', 'エリア', '規模', '決裁単位', 'HP URL', '問い合わせフォームURL 無ければメール',
               '担当部署', '相性 女性', '相性 親子', '相性 ファミリー', '相性 観光客', '相性 社員',
               '過去のWS POPUP 体験イベント実績', '実績URL', '季節の催事予定', '企業理念キーワード', '刺さる案件種類',
               '使うメール型', '優先度スコア', 'ステータス']
EXTRA_COLS = ['フォーム判定', '営業お断りの表示', 'お断り文言の抜き出し', 'メールアドレス', '判定メモ', '取得元', '地図の元データ', '取得日時', 'HP概要', '書き出し', 'HPの照合', 'LINE公式']
MAIL_TYPE = {'メーカー': 'A6', '商業施設': 'A1', '百貨店': 'A1', 'ホテル・旅館': 'A2', '観光施設': 'A2', '住宅・不動産': 'A3',
             'ブライダル': 'A4', '結婚相談所': 'A5', '自動車・保険': 'A8', '携帯ショップ': 'A8', '美容・コスメ': 'A6',
             'アパレル': 'A6', 'ライフスタイル雑貨': 'A6', '食品・生産者直売': 'A7', '教育機関': 'A10', '福祉施設': 'A10',
             '自治体・地域団体': 'A11', '企業の社内イベント': 'A9', 'その他': 'A12'}
AUDIT_COLS = ['比較電話', '比較住所', '比較値取得元', '比較値取得元URL', '比較値取得時刻']
EXTRA_COLS += ['source_detail', 'duplicate_removed', 'source_group',
               'contact_confidence', 'sales_permission', 'resume_state',
               'official_research_state', 'official_candidate_url', 'registered_url'] + AUDIT_COLS
MAIL_TYPE.update({n: d["mail"] for n, d in industries.INDUSTRIES.items()})  # 新しい業種名(広い分類)。以前の業種名の行も、上の表で引ける
INDUSTRIES = list(osm_source.INDUSTRY_TAGS)


# Google は業種名ではなく、地図で引ける言葉で探す。業種ごとの語は industries.py(広い分類ごとに複数)。
# 以前の業種名(サロン・携帯ショップなど)は、新しい分類の名前にして探す。一覧に無い言葉は「・」で分けた語をそれぞれ探す
NO_MAP = set()  # 地図の検索語にならない業種。広い分類はすべて検索語を持つ


from plan import MAP_TERMS, map_terms, estimate_minutes, plan_searches  # noqa: E402,F401  (見積もりは軽い plan.py に置く。画面が formlist を読み込まずに答えるため)
GOOGLE_PLACES_TEXT_BLOCK = "Google Places Text Search is not connected. API terms, billing, and storage/export approval are unverified."


def open_ng_state_store(db_path):
    if not db_path:
        return None
    store = NgStateStore(db_path)
    store.migration_status()
    return store


def is_ng_url(url, store=None):
    return store.is_ng([url]) if store is not None else (ng_list.is_ng(url) or ng_list.is_ng_form(url))


def add_auto_ng(url, reason, company, store=None):
    if store is None:
        return ng_list.add(url, reason, company)
    host = ng_list.domain(url)
    if not host or "." not in host or host in ng_list.TOO_WIDE:
        return False
    if ng_list._shared(host):
        u = str(url or "").strip()
        u = u if u.lower().startswith(("http://", "https://")) else "https://" + u.lstrip("/")
        return bool(store.add_form_ng(u, reason, company))
    return bool(store.add_domain_ng(url, reason, company))


def google_places_text_preview(area, term, api_key, transport, max_pages=3, request_budget=3):
    config = google_places_text_source.PlacesTextConfig(
        area=area,
        term=term,
        api_key=api_key,
        max_pages=max_pages,
        request_budget=request_budget,
    )
    result = google_places_text_source.search_text(config, transport)
    fetched_at = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = [{
        "name": p.name,
        "industry": term,
        "area": area,
        "url": p.website_uri,
        "source": "Google Places Text Search",
        "source_key": "google_places_text",
        "source_detail": f"status={result.status}; reason={result.reason}; pages={result.pages}; requests={result.requests}",
        "osm_ref": p.place_id,
        "place_id": p.place_id,
        "google_place_id": p.place_id,
        "tel": p.phone,
        "address": p.address,
        "comparison_phone": p.phone,
        "comparison_address": p.address,
        "comparison_source": "Google Places Text Search" if (p.phone or p.address) else "",
        "comparison_source_url": p.place_id,
        "comparison_fetched_at": fetched_at if (p.phone or p.address) else "",
    } for p in result.rows]
    return {"status": result.status, "reason": result.reason, "pages": result.pages, "requests": result.requests,
            "next_page_token": result.next_page_token, "rows": rows}


def reject_google_places_text_rows(rows):
    if any(r.get("source_key") == "google_places_text" for r in rows):
        raise RuntimeError(GOOGLE_PLACES_TEXT_BLOCK)


def positive_int(raw):
    try:
        value = int(raw)
    except (TypeError, ValueError):
        raise argparse.ArgumentTypeError("must be a positive integer")
    if value <= 0:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return value



_LOG_LOCK = Lock()


def log(msg):
    with _LOG_LOCK:
        print(msg, flush=True)


def ask_interactive():
    print("\n問い合わせフォームのある会社リストを作ります。\n")
    area = ""
    while not area:
        area = input("1. 探す場所を入れてください（例: 東京都世田谷区 / 大阪市北区 / 埼玉県川口市）\n   > ").strip()
    print("\n2. 業種の番号を選んでください。いくつでも選べます（例: 1 3 5）")
    for i, name in enumerate(INDUSTRIES, 1):
        print(f"   {i:>2}. {name}")
    picks = []
    while not picks:
        raw = input("   > ").replace("、", " ").replace(",", " ").split()
        picks = [INDUSTRIES[int(x) - 1] for x in raw if x.isdigit() and 1 <= int(x) <= len(INDUSTRIES)]
    raw = input("\n3. 最大何社まで調べますか（何も入れずにEnterで200）\n   > ").strip()
    limit = int(raw) if raw.isdigit() and int(raw) > 0 else 200
    return area, picks, limit


def read_input_file(path):
    rows = []
    p = Path(path)
    text = p.read_text(encoding="utf-8-sig", errors="replace") if p.suffix.lower() != ".xlsx" else ""
    for line in text.splitlines():
        cells = [c.strip() for c in next(csv.reader([line]))] if line.strip() else []
        if not cells:
            continue
        url = next((c for c in cells if c.startswith("http") or ("." in c and " " not in c and "/" in c)), "")
        name = next((c for c in cells if c and c != url), url)
        rows.append({"name": name, "industry": "", "area": "", "url": url, "source": "貼り付けリスト", "osm_ref": ""})
    return rows


def _same_store(a, b):
    """同じホームページの2行が同じ店か。同じ電話、または同じ名前と住所なら同じ店。
    電話か住所が両方にあって食い違えば別の店(本部のHPを共有する別店舗)。見分ける材料が片方に無いときは同じ店とみなす(今までどおり畳む)"""
    ta, tb = (hp_verify.tel_digits(r.get("tel", "") or r.get("phone", "")) for r in (a, b))
    if ta and ta == tb:
        return True
    aa, ab = (hp_verify.addr_key(r.get("address", "")) or hp_verify._addr_norm(r.get("address", "")) for r in (a, b))
    if aa and ab:
        name = lambda r: re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(r.get("name", "")))).lower()
        return aa == ab and name(a) == name(b)
    return not (ta and tb)


def dedupe(rows):
    """同じホームページの行のうち、同じ店の行だけを1行に畳む(畳んだ行は duplicate_removed に記録)。
    本部のHPを共有する別店舗は畳まない(同じフォームに届くかの扱いは送信側の判断)"""
    seen, out, no_url = {}, [], []
    for r in rows:
        if not r["url"]:
            no_url.append(r)
            continue
        same = seen.setdefault(canonical_url_key(r), [])
        kept = next((k for k in same if _same_store(k, r)), None)
        if kept:
            dup = f'{r.get("name", "")} / {r.get("source", "")}'.strip(" /")
            old = kept.get("duplicate_removed", "")
            kept["duplicate_removed"] = (old + " | " if old else "") + dup
            continue
        same.append(r)
        out.append(r)
    return out, no_url


def apply_ng(rows, path):
    """NGリストのドメイン(サブドメイン含む)に一致する会社を除く。対照: NGに入れたドメインの行が消える"""
    items = [l.split("#")[0].strip() for l in Path(path).read_text(encoding="utf-8-sig").splitlines()]
    ng = {ng_list.domain(x) for x in items if x and not ng_list.is_email(x) and not x.lower().startswith(ng_list.MAIL_PREFIX)} - {""}  # メールアドレスの行はドメインに化かさない(この段階では連絡先が分からない。当てるのは連絡先が判明した後の照合)
    keep = [r for r in rows if not ng_list.is_ng(r["url"], (ng, set()))]
    return keep, len(rows) - len(keep)


def judge_hp(r, res):
    """このHPがその会社のものかの結論を r["ev_text"](列「HPの照合」)に書く。
    検索で見つけたHP(r["hp_ev"]がある行)は、チェック時の判定が使えない証拠なら r["discard"] を立てる
    (別会社のHPと、そこから取ったフォーム・メールは使わない)。検索時の結果でチェック時の判定を上書きしない。
    取得元の「補完」の印は外さない(送信の関門は変えない)"""
    ev, why = res.get("ev", ""), res.get("ev_why", "")
    searched = "hp_ev" in r
    if searched and not hp_verify.usable(r, ev):
        r["discard"] = f"{ev}: {why}"
    if searched:
        r["ev_text"] = f"検索で発見・{ev}：{why}" if ev else "検索で発見・照合なし"
    else:
        r["ev_text"] = "地図に登録のHP（" + (ev or "照合なし") + (f"：{why}" if ev == "不一致" else "") + "）"
        if not ev and res.get("verdict") == "フォームあり":  # 照合の材料が取れなかったのにフォームがある行も、確認なしでは送らない
            r["source"] += " + 地図のHPは照合できず(要確認・補完扱い)"
        if ev == "不一致":  # 電話・住所の有無を問わない。地図に載っているHPでも、照合で食い違えば確認なしでは送らない
            r["source"] += " + 地図のHPが照合不一致(要確認・補完扱い)"  # 「補完」を含めて、確認なしの自動送信の関門で止める


def classify_contact(res):
    verdict = res.get("verdict", "")
    note = res.get("note", "")
    state_label = res.get("official_research_state", "")
    if state_label == "manual_candidate":
        return {"contact_confidence": "公式HP要目視", "contact_candidate": True, "sales_permission": "未確認"}
    if state_label == "not_found":
        return {"contact_confidence": "公式HP再探索済み:候補なし", "contact_candidate": False, "sales_permission": "未確認"}
    if state_label == "blocked":
        return {"contact_confidence": "公式HP再探索無効", "contact_candidate": False, "sales_permission": "未確認"}
    if state_label == "failed":
        return {"contact_confidence": "公式HP再探索失敗", "contact_candidate": False, "sales_permission": "未確認"}
    if res.get("needs_official_search"):
        confidence = "公式HP再探索待ち"
        candidate = False
    elif verdict == "フォームあり":
        confidence = "自動確認:フォーム"
        candidate = True
    elif verdict == "メールのみ":
        confidence = "自動確認:メール"
        candidate = True
    elif verdict == "要目視":
        confidence = "候補:要目視"
        candidate = True
    elif verdict in ("取得失敗", "robots禁止"):
        confidence = "取得失敗"
        candidate = False
    else:
        confidence = "未発見"
        candidate = False
    if res.get("refuse"):
        sales = "営業不可候補:拒否文あり"
    elif "問い合わせ以外" in note or "予約" in note or "採用" in note or "応募" in note:
        sales = "営業不可候補:用途不一致"
    else:
        sales = "未確認"
    return {"contact_confidence": confidence, "contact_candidate": candidate, "sales_permission": sales}


def apply_result_state(r, res, checkpoint_status=""):
    for key, val in {
        "verdict": "取得失敗",
        "form_url": "",
        "email": "",
        "refuse": False,
        "refuse_text": "",
        "note": "",
        "overview": "",
        "ev": "",
        "ev_why": "",
    }.items():
        res.setdefault(key, val)
    res.update(classify_contact(res))
    res["resume_state"] = checkpoint_status or res.get("checkpoint_status", "")
    detail = _verify_detail_text(res.get("verify_detail"))
    if detail and detail not in r.get("source_detail", ""):
        append_detail(r, detail)
    if res.get("needs_official_search"):
        detail = r.get("source_detail", "")
        r["source_detail"] = (detail + " | " if detail else "") + "official_site_search_needed"
    return res


def record_comparison_audit(r):
    """Keep independent pre-HP comparison values for later audit.

    The values come from the candidate source row (Google Maps, an injected list, etc.).
    Do not fill these from the candidate HP body, because that would make the HP check circular.
    """
    phone = r.get("comparison_phone", r.get("tel", ""))
    address = r.get("comparison_address", r.get("address", ""))
    if phone or address:
        r.setdefault("comparison_phone", phone)
        r.setdefault("comparison_address", address)
        r.setdefault("comparison_source", r.get("source") or "不明")
        r.setdefault("comparison_source_url", r.get("osm_ref", ""))
        r.setdefault("comparison_fetched_at", r.get("fetched_at", ""))
    else:
        for key in ("comparison_phone", "comparison_address", "comparison_source",
                    "comparison_source_url", "comparison_fetched_at"):
            r.setdefault(key, "")


def append_detail(r, detail):
    old = r.get("source_detail", "")
    r["source_detail"] = (old + " | " if old else "") + detail


def _verify_detail_text(detail):
    if not detail:
        return ""
    pairs = [
        ("candidate_url", detail.get("candidate_url", "")),
        ("category", detail.get("category", "")),
        ("level", detail.get("level", "")),
        ("why", detail.get("why", "")),
        ("pre_level", detail.get("pre_override_level", "")),
        ("pre_why", detail.get("pre_override_why", "")),
        ("tel_hit", detail.get("tel_hit", "")),
        ("addr", detail.get("addr_status", "")),
        ("addr_why", detail.get("addr_why", "")),
        ("title_match", detail.get("title_match", "")),
        ("industry_ok", detail.get("industry_ok", "")),
    ]
    clean = []
    for key, val in pairs:
        if val in ("", None):
            continue
        if isinstance(val, bool):
            val = str(val).lower()
        clean.append(f"{key}={str(val).replace('|', '/').replace(';', ',')[:120]}")
    return "hp_verify:" + ";".join(clean) if clean else ""


def _append_candidate_details(r):
    for detail in r.pop("_candidate_verify_details", []) or []:
        text = _verify_detail_text(detail)
        if text:
            append_detail(r, text)
    r.pop("_candidate_verify_detail", None)


def _clear_budget_stop_state(r):
    r.pop("_budget_stop_stage", None)
    r.pop("_budget_stop_current_url", None)
    r.pop("_budget_stop_candidate_url", None)
    r.pop("_budget_stop_reason", None)
    r.pop("_budget_stop_saved_at", None)


def _host(url):
    try:
        return urlparse(url if str(url).startswith("http") else "http://" + str(url)).netloc.lower().removeprefix("www.")
    except Exception:
        return ""


def row_fingerprint(r):
    """サイト確認の結果を再利用してよい「同じ行」の指紋。照合は行の名前・電話・住所に対する判定なので、
    URLが同じでも指紋が違えば(名前を直した・同じHPを使う別の店舗)前回の結果を使わない"""
    parts = (_name_key(r),
             hp_verify.tel_digits(r.get("tel", "")),  # 判定(judge_detail)が見る項目と同じにする。phone は判定が見ない
             hp_verify._addr_norm(r.get("address", "")),  # 判定は生の住所で照合するので、鍵も(正規化した)生の住所
             str(r.get("industry", "")),  # 業種で業種語の確認・施設名の判定が変わる
             str(r.get("area", "")))
    return hashlib.sha1("\x1f".join(parts).encode("utf-8")).hexdigest()[:16]


def _check_site(url, r, resume_checks=False, retry_limit=2, budget=None, line_id=None):
    checkpoint_status = ""
    variant = row_fingerprint(r) + "#" + site_checker.CHECK_RULES
    check_lease_id = None
    try:
        if resume_checks:
            action, payload = state.begin_check(url, r.get("name", ""), retry_limit=retry_limit, variant=variant)
            checkpoint_status = action
            if action in ("cached", "retry_limit", "running"):
                return payload, checkpoint_status, None
            check_lease_id = payload.get("lease_id")
        res = site_checker.check(url, row=r, budget=budget, line_id=line_id)
        if resume_checks:
            state.save_check(url, r.get("name", ""), res, check_lease_id, variant=variant)
        return res, checkpoint_status, None
    except HttpBudgetExceeded:
        raise
    except Exception as e:
        if resume_checks:
            state.fail_check(url, r.get("name", ""), e, check_lease_id, variant=variant)
        return None, checkpoint_status or "failed", e


def should_research_official_site(r, res):
    if not r.get("url") or res.get("refuse"):
        return False
    if res.get("form_url") or res.get("email") or res.get("contact_candidate"):
        return False
    if res.get("needs_official_search"):
        return True
    host = _host(r.get("url", ""))
    soft_hosts = (
        "instagram.com", "facebook.com", "x.com", "twitter.com", "ameblo.jp", "note.com",
        "lit.link", "lin.ee", "hotpepper.jp", "hpr.jp", "beauty.hotpepper.jp",
        "minimodel.jp", "ekiten.jp", "google.com", "sites.google.com", "docs.google.com",
        "forms.gle", "b-merit.jp", "reserva.be", "airrsv.net",
    )
    if any(host == h or host.endswith("." + h) for h in soft_hosts) or site_checker.is_not_company_site(host):
        return True
    return not res.get("contact_candidate") and res.get("verdict") not in ("robots禁止", "取得失敗")


def research_official_site(r, res, resume_checks=False, retry_limit=2, budget=None, line_id=None, ng_store=None):
    _clear_budget_stop_state(r)
    if not should_research_official_site(r, res):
        return res, res.get("resume_state", "")
    original_url = r.get("url", "")
    append_detail(r, f"official_site_research:from={original_url}")
    try:
        url, level, why = hp_finder.search_one(r, candidate_allowed=lambda u: candidate_allowed(u, ng_store), budget=budget, line_id=line_id)
    except HttpBudgetExceeded:
        details = r.get("_candidate_verify_details", []) or []
        if details:
            r["_budget_stop_candidate_url"] = details[-1].get("candidate_url", "")
        r["_budget_stop_stage"] = "official_site_research"
        r["_budget_stop_current_url"] = ""
        _append_candidate_details(r)
        raise
    except Exception as e:
        _append_candidate_details(r)
        append_detail(r, f"official_site_research:error={repr(e)[:80]}")
        res["official_research_state"] = "failed"
        return res, res.get("resume_state", "")
    _append_candidate_details(r)
    if not url:
        append_detail(r, f"official_site_research:not_found:{why}")
        res["official_research_state"] = "not_found"
        return res, res.get("resume_state", "")
    if _host(url) == _host(original_url):
        append_detail(r, f"official_site_research:same_host:{url}")
        res["official_research_state"] = "manual_candidate"
        res["official_candidate_url"] = url
        return res, res.get("resume_state", "")
    if level not in hp_verify.STRONG:
        append_detail(r, f"official_site_research:manual_candidate={url}; level={level}; why={why}")
        res["official_research_state"] = "manual_candidate"
        res["official_candidate_url"] = url
        return res, res.get("resume_state", "")
    if is_ng_url(url, ng_store):
        append_detail(r, f"official_site_research:blocked_by_ng={url}")
        r["discard"] = f"official_research_ng: {url}"[:120]
        res["official_research_state"] = "blocked"
        return res, res.get("resume_state", "")
    r["registered_url"] = original_url
    r["url"], r["hp_ev"], r["hp_ev_why"] = url, level, why
    r["source"] += " + 公式HP再探索(要確認)"
    append_detail(r, f"official_site_research:found={url}; level={level}; why={why}")
    found_res, checkpoint_status, err = _check_site(url, r, resume_checks=resume_checks, retry_limit=retry_limit,
                                                   budget=budget, line_id=line_id)
    if err:
        if "hp_ev" in r:
            r["discard"] = f"照合後の例外: {err!r}"[:120]
        found_res = {"verdict": "取得失敗", "form_url": "", "email": "", "refuse": False, "refuse_text": "",
                     "note": repr(err)[:120], "overview": ""}
    found_res["official_research_state"] = "found"
    return found_res, checkpoint_status


def canonical_url_key(r):
    if not r.get("url"):
        return None
    p = urlparse(r["url"] if str(r["url"]).startswith("http") else "http://" + str(r["url"]))
    return p.netloc.lower().removeprefix("www."), p.path.rstrip("/")


def candidate_allowed(candidate_url, store=None):
    return not is_ng_url(candidate_url, store)


def _name_key(r):
    return "".join(hp_finder._tokens(r.get("name", ""), r.get("area", ""))) or str(r.get("name", "")).strip().lower()


def store_identity_key(r):
    place_id = r.get("place_id") or r.get("source_place_id") or r.get("google_place_id")
    if place_id:
        return "place", str(place_id)
    url_key = canonical_url_key(r)
    if not url_key:
        return None
    tel = hp_verify.tel_digits(r.get("tel", "") or r.get("phone", ""))
    addr = hp_verify.addr_key(r.get("address", ""))
    if tel and addr:
        return "url_addr_tel", url_key, addr, tel
    if tel:
        return "url_tel", url_key, tel
    if addr:
        name = _name_key(r)
        return ("url_addr_name", url_key, addr, name) if name else None
    return None


def dedupe_checked_by_url(rows):
    seen, kept, removed = {}, [], []
    for r in rows:
        key = store_identity_key(r)
        if not key:
            kept.append(r)
            continue
        if key in seen:
            base = seen[key]
            dup = f'{r.get("name", "")} / {r.get("source", "")}'.strip(" /")
            old = base.get("duplicate_removed", "")
            base["duplicate_removed"] = (old + " | " if old else "") + dup
            removed.append(r)
            continue
        seen[key] = r
        kept.append(r)
    return kept, removed


def is_output_ng(r, store=None):
    if store is not None:
        return store.is_ng([r.get("url", ""), r.get("form_url", "")])
    return (
        ng_list.is_ng(r.get("url", ""))
        or ng_list.is_ng_form(r.get("url", ""))
        or ng_list.is_ng_form(r.get("form_url", ""))
    )


def assign_budget_line_ids(rows):
    for i, r in enumerate(rows, 1):
        if not r.get("_budget_line_id"):
            r["_budget_line_id"] = f"row-{i}:{source_key(r)}:{r.get('name', '')}:{r.get('area', '')}"


BUDGET_STOP_COLS = [
    "name", "industry", "area", "hp_url", "stage", "current_url", "candidate_url", "source",
    "source_detail", "stop_reason", "saved_at",
]


def _record_budget_stop_row(r, reason, stage="", current_url=""):
    details = list(r.get("_candidate_verify_details", []) or [])
    candidate_url = r.get("_budget_stop_candidate_url", "")
    if not candidate_url and details:
        candidate_url = details[-1].get("candidate_url", "")
    stage = r.get("_budget_stop_stage", stage)
    current_url = r.get("_budget_stop_current_url", current_url)
    _append_candidate_details(r)
    if stage != "official_site_research" and "hp_verify:" not in str(r.get("source_detail", "")) and r.get("url"):
        text = _verify_detail_text({
            "candidate_url": r.get("url", ""),
            "category": "fetch_failed",
            "why": "http_budget_exceeded",
        })
        if text:
            append_detail(r, text)
    append_detail(r, f"http_budget_stop:reason={str(reason).replace('|', '/')[:120]}")
    stopped = dict(r)
    stopped["_budget_stop_reason"] = str(reason)
    stopped["_budget_stop_stage"] = stage
    stopped["_budget_stop_current_url"] = current_url
    stopped["_budget_stop_candidate_url"] = candidate_url
    stopped["_budget_stop_saved_at"] = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    return stopped


def write_budget_stop_diagnostic(rows):
    OUT_DIR.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    for suffix in [""] + [f"_{i:02d}" for i in range(1, 1000)]:
        path = OUT_DIR / f"http_budget_stop_diagnostic_{stamp}{suffix}.csv"
        try:
            with open(path, "x", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(BUDGET_STOP_COLS)
                for r in rows:
                    w.writerow(safe_output_row([
                        r.get("name", ""),
                        r.get("industry", ""),
                        r.get("area", ""),
                        r.get("url", ""),
                        r.get("_budget_stop_stage", ""),
                        r.get("_budget_stop_current_url", ""),
                        r.get("_budget_stop_candidate_url", ""),
                        r.get("source", ""),
                        r.get("source_detail", ""),
                        r.get("_budget_stop_reason", ""),
                        r.get("_budget_stop_saved_at", ""),
                    ]))
            return path
        except FileExistsError:
            continue
    raise FileExistsError(f"could not create unique http budget diagnostic file for {stamp}")


def run_checks(rows, workers=4, resume_checks=False, retry_limit=2, official_research=False, budget=None, ng_store=None, stop_check=None):
    """stop_check(止める指示を見る関数)が真になったら、まだ調べていない会社は調べずに外す(戻りに入れない。件数は RUN_STATS["not_checked_stopped"])"""
    assign_budget_line_ids(rows)
    done = [0]
    budget_stop_rows = []
    budget_stop_lock = Lock()

    def one(r):
        if stop_check and stop_check():
            with budget_stop_lock:
                RUN_STATS["not_checked_stopped"] = RUN_STATS.get("not_checked_stopped", 0) + 1
            return None
        _clear_budget_stop_state(r)
        record_comparison_audit(r)
        checkpoint_status = ""
        line_id = hp_finder._line_id(r)
        try:
            res, checkpoint_status, err = _check_site(r["url"], r, resume_checks=resume_checks, retry_limit=retry_limit,
                                                      budget=budget, line_id=line_id)
            if err:
                raise err
            res = apply_result_state(r, res, checkpoint_status)
            if official_research:
                res, checkpoint_status = research_official_site(r, res, resume_checks=resume_checks, retry_limit=retry_limit,
                                                                budget=budget, line_id=line_id, ng_store=ng_store)
                res = apply_result_state(r, res, checkpoint_status)
            judge_hp(r, res)
        except HttpBudgetExceeded as e:
            try:
                stopped = _record_budget_stop_row(r, e, stage="initial_check", current_url=r.get("url", ""))
                with budget_stop_lock:
                    budget_stop_rows.append(stopped)
            except Exception as record_error:
                log(f"HTTP budget stopped; diagnostic row record failed: {record_error!r}")
            raise
        except Exception as e:  # keep going; record the raw error
            if "hp_ev" in r:
                r["discard"] = f"照合中の例外: {e!r}"[:120]  # 検索で見つけたHPは、照合できなければ使わない
            res = {"verdict": "取得失敗", "form_url": "", "email": "", "refuse": False, "refuse_text": "", "note": repr(e)[:120], "overview": ""}
            res = apply_result_state(r, res, checkpoint_status or "failed")
        done[0] += 1
        log(f"  [{done[0]}/{len(rows)}] {res['verdict']}{' 営業お断り' if res['refuse'] else ''}  {r['name']}")
        return {**r, **res, "checked_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}

    try:
        with ThreadPoolExecutor(max_workers=workers) as ex:  # different sites in parallel, same site sequential
            return [x for x in ex.map(one, rows) if x is not None]
    except HttpBudgetExceeded:
        try:
            path = write_budget_stop_diagnostic(budget_stop_rows)
            log(f"HTTP budget stopped; diagnostic saved: {path}")
        except Exception as save_error:
            log(f"HTTP budget stopped; diagnostic save failed: {save_error!r}")
        raise


def size_text(r):
    return (r.get('size', '') + ('：' + r['size_why'] if r.get('size_why') else '')) if r.get('size') else ''


def to_row(i, r):
    contact = r["form_url"] if r["verdict"] in ("フォームあり", "要目視") else (r["email"] or r["form_url"])
    status = "対象外 営業お断り" if r["refuse"] else "未着手"
    ledger = [i, r["name"], r["industry"], r["area"], size_text(r), "", r["url"], contact, "", "", "", "", "", "", "", "", "", "", "",
              MAIL_TYPE.get(r["industry"], ""), "", status]
    extra = [r["verdict"], "あり" if r["refuse"] else "", r["refuse_text"], r["email"], r["note"], r["source"], r["osm_ref"], r["checked_at"], r.get("overview", ""), "", r.get("ev_text", ""), r.get("line_url", "")]
    extra += [r.get("source_detail", ""), r.get("duplicate_removed", ""), r.get("source_key", ""),
              r.get("contact_confidence", ""), r.get("sales_permission", ""), r.get("resume_state", ""),
              r.get("official_research_state", ""), r.get("official_candidate_url", ""), r.get("registered_url", ""),
              r.get("comparison_phone", ""), r.get("comparison_address", ""), r.get("comparison_source", ""),
              r.get("comparison_source_url", ""), r.get("comparison_fetched_at", "")]
    return ledger + extra



def safe_output_cell(value):
    if not isinstance(value, str) or not value:
        return value
    had_leading_control = ord(value[0]) < 32
    value = "".join(ch if ch in "\t\r\n" or ord(ch) >= 32 else " " for ch in value)
    trimmed = value.lstrip(" \t\r\n")
    if trimmed and trimmed[0] in ("=", "+", "-", "@"):
        return "'" + value
    if had_leading_control:
        return "'" + value
    return value


def safe_output_row(row):
    return [safe_output_cell(value) for value in row]
ORDER = {"フォームあり": 0, "要目視": 1, "メールのみ": 2, "見つからず": 3, "robots禁止": 4, "取得失敗": 5}


def source_key(r):
    if r.get("source_key"):
        return r["source_key"]
    src = r.get("source", "")
    if "Overture" in src:
        return "overture"
    if "OpenStreetMap" in src:
        return "osm"
    if "ポータル" in src:
        return "portal"
    if "リスト" in src:
        return "input"
    if "Google" in src:
        return "google"
    return src.split(":")[0] or "unknown"


def summarize_source_funnel(all_rows, deduped_targets, checked, no_url):
    keys = sorted({source_key(r) for r in all_rows + deduped_targets + checked + no_url})
    log("\n=== source funnel ===")
    log("source, candidates, hp_before_dedupe, no_hp, deduped_targets, form_or_email, usable_after_checks")
    for key in keys:
        src_rows = [r for r in all_rows if source_key(r) == key]
        src_targets = [r for r in deduped_targets if source_key(r) == key]
        src_checked = [r for r in checked if source_key(r) == key]
        contactable = sum(1 for r in src_checked if r.get("verdict") in ("フォームあり", "要手動", "メールのみ"))
        usable = sum(1 for r in src_checked if r.get("verdict") in ("フォームあり", "要手動", "メールのみ") and not r.get("refuse"))
        contactable = sum(1 for r in src_checked if r.get("contact_candidate"))
        usable = sum(1 for r in src_checked if r.get("contact_candidate") and r.get("sales_permission") == "未確認")
        log(f"{key}, {len(src_rows)}, {sum(1 for r in src_rows if r.get('url'))}, {sum(1 for r in src_rows if not r.get('url'))}, {len(src_targets)}, {contactable}, {usable}")


def summarize_contact_states(checked):
    automatic = sum(1 for r in checked if r.get("contact_confidence") in ("自動確認:フォーム", "自動確認:メール"))
    candidates = sum(1 for r in checked if r.get("contact_candidate"))
    missing = sum(1 for r in checked if r.get("contact_confidence") == "未発見")
    official_wait = sum(1 for r in checked if r.get("contact_confidence") == "公式HP再探索待ち")
    not_allowed = sum(1 for r in checked if str(r.get("sales_permission", "")).startswith("営業不可候補"))
    log("\n=== contact readiness ===")
    log(f"自動確認済み(フォーム/メール): {automatic}件")
    log(f"候補込み(要目視含む): {candidates}件")
    log(f"未発見: {missing}件")
    log(f"公式HP再探索待ち: {official_wait}件")
    log(f"営業不可候補(拒否文/用途不一致): {not_allowed}件")


def write_outputs(checked, no_url, label):
    reject_google_places_text_rows(list(checked) + list(no_url))
    OUT_DIR.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    safe = "".join(c for c in label if c not in '\\/:*?"<>|# ')[:30] or "リスト"
    base = OUT_DIR / f"フォームリスト_{safe}_{stamp}"
    n = 1
    while True:  # 同じ場所を同じ分に2つ探しても、互いの結果を上書きしない(.csv を排他で先に作って名前を確保する)
        try:
            open(base.with_suffix(".csv"), "x").close()
            break
        except FileExistsError:
            n += 1
            base = OUT_DIR / f"フォームリスト_{safe}_{stamp}_{n}"
    checked = sorted(checked, key=lambda r: (r["refuse"], ORDER.get(r["verdict"], 9)))
    rows = [safe_output_row(to_row(i, r)) for i, r in enumerate(checked, 1)]
    header = LEDGER_COLS + EXTRA_COLS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "企業リスト"
    ws.append(header)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
    gray = PatternFill("solid", fgColor="DDDDDD")
    for r_idx, r in enumerate(checked, 2):
        if r["refuse"]:
            for c in ws[r_idx]:
                c.fill = gray
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for col, w in zip("ABCDEFGH", (5, 30, 14, 14, 6, 8, 36, 44)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("URLが分からない会社")
    ws2.append(["企業名", "業種", "エリア", "検索用リンク（クリックでGoogle検索）", "地図の元データ"])
    for r in no_url:
        q = quote_plus(f'{r["name"]} {r["area"]}')
        ws2.append(safe_output_row([r["name"], r["industry"], r["area"], f"https://www.google.com/search?q={q}", r["osm_ref"]]))
    ws3 = wb.create_sheet("このリストで確かめていないこと")
    for line in NOT_CHECKED:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    wb.save(base.with_suffix(".xlsx"))
    with open(base.with_suffix(".csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return base.with_suffix(".xlsx"), base.with_suffix(".csv")


RUN_STATS = {"out_of_area": 0}
import os as _os


def summary_name(port=None):
    """複数のサーバーが同時に探すので、内訳の印はサーバー(ポート)ごとに分ける"""
    return f"最後の探索_内訳_{port or _os.environ.get('FORMLIST_PORT') or 8780}.json"


SUMMARY_FILE = summary_name()


RUN_ID = {"started": "", "id": ""}

OUTCOME_MESSAGE = {
    "running": "探している途中です。この印のまま探す処理が終わっていたら、途中で止まった(止めた・異常終了した)回です。",
    "ok": "会社が見つかりました。",
    "zero": "Googleマップの検索は失敗せず、取りこぼしの記録もありませんが、会社は0件でした。",
    "auto_stopped": "0件が続いたため自動で止めました。すべての検索を最後まで実行したわけではないので、0件とは言い切れません。場所の書き方（市区町村まで）と、業種・キーワードを変えて探し直してください。",
    "all_large": "地図では会社が見つかりましたが、すべて大手・チェーンと推定して外しました。0件ではありません。「大手・チェーンも含める」を入れて、探し直してください。",
    "all_outside": "地図では会社が見つかりましたが、住所が指定の地域と合わず、全部「エリア外」として外しました。0件ではありません。地名の書き方を確かめて、市区町村まで1つずつ「、」で区切って入れ直してください（例: 東京都江東区、東京都墨田区）。",
    "missed": "失敗はありませんが、読めなかった店・名前を読めなかった店・施設ページに直行して読めなかった検索があり、残った会社は0件でした。0件と言い切れません。時間をおいて探し直してください。",
    "failed": "取得に失敗しました。これは「0件」ではありません。通信の失敗か、Googleに止められた可能性があります。",
    "partial": "一部の検索が失敗しましたが、取れた分で続けました。取りこぼしがあります。",
    "stopped": "止める指示で、途中で止めました。調べ終わった分だけをリストにしています（リストが無い回もあります）。",
}

# このリストで確かめていないこと(結果の出力自身に書く。勝利条件10)
NOT_CHECKED = [
    "Googleマップは1回の検索で最後まで出さないことがあります。読めなかった店・上限の件数に届く前に一覧が止まった検索の数は、探した記録の内訳(map)にあります",
    "「フォームあり」は、そのホームページが会社のものかを確かめる前の判定を含みます。送ってよいかは「HPの照合」の列と、自動送信の対象かどうかで決まります",
    "予約・掲載・口コミ・SNSのページは、表にあるものだけを外しています。表に無い掲載サイトの1店ページは、会社のホームページとして残ることがあります",
    "大手・チェーンかどうかは、名前と同じホームページの店の数だけで推定しています（売上・従業員数は見ていません）",
    "エリア外かどうかは住所の文字で判定しています。住所が無い店は判定していません",
    "営業お断りの表示は、文字で書かれたものだけを見ています（画像の中の文字は見ていません）",
    "画面の描画に頼るフォーム、ログインが要るページ、確認の画像(CAPTCHA)だけのページは、フォームの有無を確かめられていません",
]


def map_counts():
    """この回の地図の検索の、取れた数と取りこぼした数(map_browser_source が数えたもの)"""
    c = dict(map_browser_source.RUN_COUNTS)
    return {"searches": c.get("searches", 0), "rows": c.get("rows", 0), "unreadable": c.get("unreadable", 0),
            "cut_short": c.get("cut_short", 0), "place_missed": c.get("place_missed", 0), "place_read": c.get("place_read", 0),
            "empty": c.get("empty", 0), "bad_name": c.get("bad_name", 0)}


def missed_total(m=None):
    m = m or map_counts()
    return m["unreadable"] + m["place_missed"] + m["bad_name"]  # cut_short(一覧が上限前に止まった)は、店がそれだけしか無い場合と見分けられないので数えない


def build_funnel(fetched, no_url_first, with_url, large_excluded, cut, examined, dropped, refused, checked):
    """この回の段階別の件数(どこで何社が落ちたか)。画面の「結果を見る」に出す。各行は別々に数えた件数で、足し引きで合うとは限らない
    (同じ店が複数の理由で外れることがある)。label と n と、読み方の note"""
    counts = {}
    for r in checked:
        counts[r.get("verdict") or "未判定"] = counts.get(r.get("verdict") or "未判定", 0) + 1
    f = [
        {"label": "地図などで取れた店・施設", "n": fetched, "note": f"エリア外として外した {RUN_STATS.get('out_of_area', 0)}件、同じ店の重複として1件にまとめた {RUN_STATS.get('duplicate_places', 0)}件は、この数に入っていません"},
        {"label": "ホームページが分かった", "n": with_url, "note": "同じホームページの店は1社にまとめたあとの数"},
        {"label": "ホームページが分からない", "n": no_url_first, "note": "名前での検索でも、電話・住所が合うホームページを見つけられなかった店"},
    ]
    if large_excluded:
        f.append({"label": "大手・チェーンとして外した", "n": large_excluded, "note": "「大手・チェーンも含める」を入れると、リストに残ります"})
    if cut:
        f.append({"label": "上限を超えて調べなかった", "n": cut, "note": "「上限の社数」を増やすと調べます"})
    f.append({"label": "ホームページを開いて調べた", "n": examined, "note": ""})
    if dropped:
        f.append({"label": "別会社のページの疑いで外した", "n": dropped, "note": "検索で見つけたホームページが、別の会社のものだった疑い"})
    if refused:
        f.append({"label": "営業お断りの表示があった", "n": refused, "note": "送らないリストに入れ、一覧から外しました"})
    for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
        f.append({"label": f"判定: {k}", "n": v, "note": ""})
    return f


def write_run_summary(outcome, **extra):
    """この回の探索の結果を、MCP・画面が読めるJSONで残す。0件(本当に無かった)と失敗(止められた・通信の失敗)を分けるための印。
    outcome: running(探している途中) / ok(会社が見つかった) / zero(失敗も取りこぼしも無く0件) / all_outside(取れたが全部エリア外) / all_large(取れたが全部大手・チェーンとして外した) /
             missed(失敗は無いが取りこぼしがあり0件) / failed(失敗があり、取れた会社が無い) / partial(失敗があり、取れた分だけで続けた) / stopped(止める指示で止めた)
    run_id・started はこの回の印。読む側は、探す処理を始めた時刻より古い started の印を、この回のものとして使わない"""
    import json
    warnings = list(map_browser_source.RUN_WARNINGS)  # やり直しても半数以上読めなかった検索語。MCP の結果が読む not_checked の最後にも足す
    data = {"outcome": outcome, "message": OUTCOME_MESSAGE.get(outcome, ""), "run_id": RUN_ID["id"], "started": RUN_ID["started"],
            "finished": dt.datetime.now().isoformat(timespec="seconds") if outcome != "running" else "",
            "out_of_area": RUN_STATS.get("out_of_area", 0), "fetch_failed": RUN_STATS.get("fetch_failed", 0),
            "duplicate_places": RUN_STATS.get("duplicate_places", 0), "map": map_counts(), "missed": missed_total(),
            "area_notes": RUN_STATS.get("area_notes", []), "not_checked": NOT_CHECKED + warnings, "term_warnings": warnings, **extra}
    if warnings and outcome != "running":  # 画面のログの最後に出す(途中のログに埋もれさせない)
        log("\n[警告] 読めなかった店が多い検索語があります:\n" + "\n".join("  " + w for w in warnings))
    try:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        import filelock
        path = OUT_DIR / summary_name()
        with filelock.locked(path):
            filelock.atomic_write(path, json.dumps(data, ensure_ascii=False, indent=2))
    except Exception as e:  # 印が書けなくても、収集の結果そのものは壊さない
        log(f"（探索の内訳の印を書けませんでした: {e!r}）")


def out_of_area(r, area):
    """行の住所が、指定した区・市の外にある(Googleが近くの別の区の店も返す)。住所が無い・エリアに区市町村が無いときは判定しない(False)。
    住所に都道府県があれば、それも指定と合うかを見る(東京都港区の指定で、大阪市港区の店を残さない)。ヶ/ケ・ヵ/カは同じ字として比べる"""
    scopes = subareas.area_scopes(area)
    if not scopes or not r.get("address"):
        return False
    return not subareas.address_in(r["address"], scopes)


def dedupe_places(rows):
    """町名・検索語をまたいで同じ店が何回も出たものを、1行にまとめる(Googleマップは町名で探しても近くの店を返す)。
    鍵は施設の識別子(place_id)、無ければ 名前+ホームページ+住所。大手・チェーンの推定と件数の集計より前に通す
    (2026-10-21 破壊レポート 致命傷2: 豊洲・東雲・有明の3検索に出た1店が「同じホームページが3件」でチェーン扱いになっていた)。
    戻り: (まとめた後の行, まとめて消した件数)"""
    seen, out = set(), []
    for r in rows:
        pid = r.get("place_id") or r.get("google_place_id")
        if not pid and r.get("source") != "Googleマップ":
            out.append(r)  # 貼り付けリスト・ポータル・OSM の行はここでは触らない(同じHPの重複は dedupe が記録つきで外す)
            continue
        if pid:
            key = ("id", str(pid))
        else:
            key = ("name", unicodedata.normalize("NFKC", str(r.get("name", ""))).strip().lower(), str(r.get("url", "")).strip().lower().rstrip("/"),
                   hp_verify._addr_norm(r.get("address", "")))
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out, len(rows) - len(out)


def limit_osm_rows(rows, limit, log):
    """OSM の取得件数を --limit で切る。電話も住所も無い行が半数以上なら、HPを検索で補えず照合が通りにくいと警告する"""
    no_evidence = sum(1 for r in rows if not r.get("tel") and not r.get("address"))
    if rows and no_evidence * 2 >= len(rows):
        log(f"  [osm] 電話も住所も無い行が {no_evidence}/{len(rows)}件あります。HPを検索で補っても、電話・住所で照合できず採れません（取得元 google の方が、電話・住所を持つ行が多いです）")
    if len(rows) > limit:
        log(f"  [osm] {len(rows)}件のうち、{limit}件だけを使います（--limit。HP・電話・住所を持つ行を先に残します）")
        return sorted(rows, key=lambda r: (not r.get("url"), not (r.get("tel") or r.get("address"))))[:limit]
    return rows


def count_stages(fetched, checked, no_url, skip_reasons):
    """段階別の件数を辞書で返す(出力の組み立てと検算を分けるため)。
    候補 = HP不明 + 調べた + 調べなかった(同じHP・大手チェーン・上限超え・お断り表示で外した等。引き算で出す)
    調べた = 自動送信の対象 + 対象にならない(理由別の合計)"""
    reasons = {}
    for _name, _url, reason in skip_reasons:
        reasons[reason] = reasons.get(reason, 0) + 1
    not_target = sum(reasons.values())
    return {
        "fetched": fetched,
        "no_url": len(no_url),
        "checked": len(checked),
        "not_checked": fetched - len(no_url) - len(checked),
        "reasons": reasons,
        "sendable": len(checked) - not_target,
    }


def summarize_sendable(csv_path, fetched, checked, no_url):
    """出力の最後に、候補から「自動送信の対象」までの段階別の件数と、落ちた理由別の件数を出す。
    自動送信の対象の判定は autosend と同じ関数(target_skip_reasons)を使う(別の数え方をしない)"""
    import autosend  # 遅延: autosend が formlist を import するため
    outside = RUN_STATS.get("out_of_area", 0)
    skips = autosend.target_skip_reasons(csv_path)
    c = count_stages(fetched, checked, no_url, skips)  # fetched はエリア外を除いた後の数
    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        csv_rows = list(csv.DictReader(fh))
    verified = [r for r in csv_rows if not autosend._hp_verify_skip_reason(r)]
    form_ok = [r for r in verified if r.get(EXTRA_COLS[0]) == next(iter(ORDER))]
    dropped = sum(1 for r in no_url if r.get("discard"))
    c.update(verified=len(verified), form_ok=len(form_ok), dropped=dropped, outside=outside)
    mail_skips = autosend.mail_target_skip_reasons(csv_path)
    mail_only = sum(1 for r in csv_rows if r.get(EXTRA_COLS[0]) == "メールのみ")
    mail_reasons = {}
    for _n, _m, why in mail_skips:
        mail_reasons[why.split(":")[0]] = mail_reasons.get(why.split(":")[0], 0) + 1
    c.update(mail_only=mail_only, mail_sendable=mail_only - len(mail_skips), mail_reasons=mail_reasons)
    log("\n=== 件数の内訳(候補から自動送信の対象まで) ===")
    log(f"地図・リストの候補: {c['fetched'] + outside}件")
    if outside:
        log(f"  住所が指定エリアの外で除外: {outside}件")
    log(f"  ホームページ不明: {c['no_url'] - dropped}件 / 検索で見つけたHPが別会社の疑いで捨てた: {dropped}件")
    log(f"  調べなかった(同じHP・NG・大手チェーン・上限超え・お断り表示で外した等): {c['not_checked']}件")
    log(f"  調べた: {c['checked']}件")
    log(f"    うちHPの照合が通った(電話一致・住所一致): {c['verified']}件 / 照合で落ちた・照合できなかった: {c['checked'] - c['verified']}件")
    log(f"    照合が通った会社のうち、フォームあり: {c['form_ok']}件")
    grouped = {}
    for reason, n in c["reasons"].items():
        key = reason.split(":")[0]
        grouped[key] = grouped.get(key, 0) + n
    not_form = "フォーム判定がフォームありではない"
    if not_form in grouped:  # メールのみの会社は、下の「メールのみの会社」で別に数える(フォームが無いというだけの一括りにしない)
        grouped[not_form] -= c["mail_only"]
        if grouped[not_form] <= 0:
            del grouped[not_form]
    for reason, n in sorted(grouped.items(), key=lambda kv: -kv[1]):
        label = "フォームもメールのみでもない" if reason == not_form else reason
        log(f"    自動送信の対象にならない({label}): {n}件")
    log(f"  自動送信の対象: {c['sendable']}件（調べた{c['checked']}件中）")
    log(f"  メールのみの会社: {c['mail_only']}件 → メール対象 {c['mail_sendable']}件 / 対象外 {c['mail_only'] - c['mail_sendable']}件")
    for reason, n in sorted(c["mail_reasons"].items(), key=lambda kv: -kv[1]):
        log(f"    メール対象にならない({reason}): {n}件")
    if c["mail_only"]:
        log("  ※メールの送信機能はまだありません(第2段で追加)。いまは、メール対象の件数を確かめるだけで、自動では送りません")
    line_only = [r for r in csv_rows if r.get("LINE公式") and r.get(EXTRA_COLS[0]) != next(iter(ORDER))]
    c["line_only"] = len(line_only)
    log(f"  フォームが無くLINE公式がある会社: {c['line_only']}件（自動では送らない。列「LINE公式」のURLを開き、友だち追加して手で送る）")
    log("  ※「フォームあり」は、そのHPが会社のものかの確認前の数です。送れる件数は、この「自動送信の対象」です")
    return c


def summarize(fetched, checked, no_url, refused=0):
    counts = {}
    for r in checked:
        counts[r["verdict"]] = counts.get(r["verdict"], 0) + 1
    refuse = refused
    log("\n=== 結果 ===")
    log(f"地図・リストから見つかった会社: {fetched}件")
    log(f"  ホームページが分かった会社（同じHPの重複を除く）: {len(checked)}件を調べました / ホームページ不明: {len(no_url)}件")
    sup = sum("補完" in r["source"] for r in checked)
    if sup:
        log(f"  うち {sup}件は検索で見つけたHPです（チェーン店は本部のHP。取得元の列に「補完」と入っています）")
    for k in ORDER:
        if counts.get(k):
            log(f"  {k}（HPの照合前の数）: {counts[k]}件")
    log(f"  営業お断りの表示あり（送らない）: {refuse}件")


from plan import PARALLEL_MAX  # noqa: E402  (Googleマップへの同時アクセスの上限。増やすほどブロックされやすい)
DEFAULT_FETCH_LIMIT = 200  # 検索する数の上限の既定値(候補を集める全体の件数)。出力する数(--limit)とは連動させない(連動させると出力の指定に合わせて集める側も膨らんで終わらなくなるため)
BLOCK_MINUTES = 30  # Googleに止められたら、この時間はどのサーバーも新しい地図の検索を始めない(暫定値。止められた時間の実測は無い)
TRANSIENT_IN_A_ROW = 3  # 地図の結果が出ない失敗がこの回数続いたら、制限の可能性があるとして残りの検索を止める
from plan import SECONDS_PER_CARD  # noqa: E402
# 次の2つは暫定値(2026-10-26 設定。0件が続いたときに止める仕組み。実測の根拠は無い。90_作業メモ/業種検索語の検証/ の結果で見直す)
ZERO_TERM_IN_A_ROW = 6  # 同じ検索語が、町名を変えて続けて0件(1件も引っかかっていない)なら、その検索語の残りは飛ばす
# 全体を止めるのは、すべての検索語が「町名を変えても1件も引っかからない」(上の6回)状態になったとき。検索語をまたいで0件を数えない(あとの検索語が実行されなくなるため)


class NotRun(RuntimeError):
    """止める指示・0件が続いたための自動停止で、実行しなかった検索。取得の失敗ではない"""


class BlockedElsewhere(map_browser_source.GoogleBlocked):
    """同時に動いている別のサーバーがGoogleに止められた印を見て、この検索を始めなかった(印を書き直して時間を延ばさない)"""


def block_file():
    return OUT_DIR / "google_止められた.json"


def mark_google_blocked(reason):
    """Googleに止められた印を、出力フォルダに残す。同時に動いている別のサーバーも、これを読んで新しい検索を始めない"""
    import json
    import filelock
    now = dt.datetime.now()
    data = {"since": now.isoformat(timespec="seconds"), "until": (now + dt.timedelta(minutes=BLOCK_MINUTES)).isoformat(timespec="seconds"),
            "port": _os.environ.get("FORMLIST_PORT") or "8780", "reason": str(reason)[:200]}
    try:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        with filelock.locked(block_file()):
            filelock.atomic_write(block_file(), json.dumps(data, ensure_ascii=False, indent=2))
    except Exception as e:  # 印が書けなくても、この回の止め方は変えない
        log(f"（Googleに止められた印を書けませんでした: {e!r}）")


def google_block_active():
    """Googleに止められた印がまだ有効なら、その終わる時刻の文字列。無ければ空"""
    import json
    try:
        d = json.loads(block_file().read_text(encoding="utf-8-sig"))
        until = dt.datetime.fromisoformat(d["until"])
    except (OSError, ValueError, KeyError, TypeError):
        return ""
    return d["until"] if until > dt.datetime.now() else ""


def stop_file(port=None):
    """止める指示の印。MCP・画面がこのファイルを置くと、探す処理は次の区切り(検索1回・会社1社)で止まる"""
    return OUT_DIR / f"探す_止める_{port or _os.environ.get('FORMLIST_PORT') or 8780}.flag"


def stop_requested():
    try:
        return stop_file().exists()
    except OSError:
        return False


def clear_stop():
    try:
        stop_file().unlink()
    except OSError:
        pass




def search_plan_message(budget, n_jobs, minutes, workers):
    """探す前に出す説明。実際に走る検索の上限と打ち切りの挙動を先に、全部の検索の数と所要時間は参考値として後ろに置く"""
    return (f"\n実際に走る検索の上限: 候補が全体で{budget}件（検索する数の上限）に達したら、残りの検索は始めずに次へ進みます。"
            f"候補が0件の検索はそこで終えて次の検索へ進み、0件が続くと自動で止まります。画面の「ここで止める」でもいつでも止められます。"
            f"（参考）全部の検索（検索語×町名）は最大{n_jobs}件で、最後まで走った場合の地図だけの所要時間は最大およそ{minutes}分です。"
            f"上限に達すればそれより早く終わります。最大{workers}本ずつ並行して取得し、1件ごとに「[終わった数/全体]」が出ます。このあとホームページの検索とフォームの確認が続きます。")


def prefetch_google(area, units, limit, workers, block_check=None, stop_check=None, fetch_limit=None):
    """Googleマップの検索を、検索語(業種・キーワード)×町名の単位で並列に取る。戻りは {(単位の番号, 検索語, 町): 行の一覧 か 例外}。
    どれか1つがGoogleに止められたら、まだ始めていない検索は実行しない(ブロックされたときに叩き続けないため)。
    地図の結果が出ない失敗が TRANSIENT_IN_A_ROW 回続いたときも、制限の可能性として止める。
    block_check: 別のサーバーがGoogleに止められた印を見る関数(有効なら真)。stop_check: 止める指示を見る関数。
    同じ検索語×町名は1回だけ取る(業種「サロン」と「美容・コスメ」の両方にある「美容室」など)。戻りの行は単位ごとに別の辞書の写し。
    ブラウザは1検索ごとに自前で開いて閉じるので、スレッド間で共有するものは無い"""
    import random
    import time
    from threading import Event
    stop = Event()
    lock = Lock()
    streak = [0]
    subs = subareas.split(area)
    keys = []
    for ui, (ind, kw) in enumerate(units):
        if not kw and ind in NO_MAP:
            continue
        for term in ([kw] if kw else map_terms(ind)):
            for sub in subs:
                keys.append((ui, term, sub))
    jobs = list(dict.fromkeys((term, sub) for _ui, term, sub in keys))
    prog = {"done": 0}  # 終わった検索の数
    all_terms = {t for t, _s in jobs}
    term_hit, term_zero = {}, {}  # 検索語ごとの、1件でも引っかかった回数 / 続けて0件だった回数
    auto_stop = []  # 0件が続いて自動で止めた理由(1件だけ入る)
    budget = max(fetch_limit, 1) if fetch_limit else DEFAULT_FETCH_LIMIT  # 検索する数の上限(候補を取得した行の数)。集める段階の停止条件はこれ1つだけ(出力上限 limit とは連動させない)。上限が無いと、検索語×町名の数だけ際限なく開き続ける(2026-10-28 実例: 50社の指定で120検索・約3,600件の読み取りが半日続いた)
    fetched_total = [0]  # ここまでに取得した行の数
    budget_stop = []  # 候補の上限に達して自動で止めた理由(1件だけ入る)

    def finish(job, n, note=""):
        """1検索が終わるたびに、どの検索が終わったか・何件目かを出す(同じ文言が並んで止まって見えないように)"""
        with lock:
            prog["done"] += 1
            k = prog["done"]
        log(f"  [{k}/{len(jobs)}] {job[0]} {job[1]}: {n}件{note}")

    def run(job):
        term, sub = job
        if stop.is_set():
            if auto_stop:
                finish(job, 0, "（0件が続いたため実行せず）")
                return job, NotRun(auto_stop[0])
            return job, RuntimeError("ほかの検索が止まったため、この検索は実行しませんでした")
        if stop_check and stop_check():
            stop.set()
            finish(job, 0, "（止める指示のため実行せず）")
            return job, NotRun("止める指示があったため、この検索は実行しませんでした")
        if block_check and block_check():
            stop.set()
            return job, BlockedElsewhere("同時に動いている別の検索がGoogleに止められたため、この検索は実行しませんでした。時間をおいて探し直してください")
        with lock:
            skip = term_zero.get(term, 0) >= ZERO_TERM_IN_A_ROW and not term_hit.get(term)
        if skip:  # この検索語は、町名を変えても1件も引っかからない。残りの町名は飛ばす
            with lock:
                RUN_STATS.setdefault("zero_terms_skipped", {}).setdefault(term, 0)
                RUN_STATS["zero_terms_skipped"][term] += 1
            finish(job, 0, f"（「{term}」は{ZERO_TERM_IN_A_ROW}町続けて0件のため飛ばしました）")
            return job, []
        with lock:
            over_budget = fetched_total[0] >= budget
            first_budget = over_budget and not budget_stop
            if first_budget:
                budget_stop.append(f"候補の取得が全体の上限（{budget}件）に達したため、残りの検索は自動で止めました。取れた分でリストを作ります")
                RUN_STATS["auto_stopped_budget"] = budget_stop[0]
        if over_budget:  # 候補はもう十分に集まっている。これ以上は開かない
            if first_budget:
                log(f"\n  {budget_stop[0]}")
            finish(job, 0, "（候補の上限に達したため実行せず）")
            return job, NotRun(budget_stop[0])
        time.sleep(random.uniform(0, 2))  # 同時に始めた検索が、同じ瞬間にGoogleへ集中しないようにずらす
        try:
            got = map_browser_source.fetch("google", sub, term, log, limit=limit)
        except map_browser_source.GoogleBlocked as e:
            stop.set()  # 確認画面・回数制限。まだ始めていない検索は実行しない
            return job, e
        except Exception as e:
            with lock:
                streak[0] += 1
                if streak[0] >= TRANSIENT_IN_A_ROW and not stop.is_set():
                    stop.set()
                    log(f"  地図の結果が出ない失敗が{streak[0]}回続きました。Googleの制限の可能性があるため、残りの検索は始めません")
            finish(job, 0, f"（失敗: {type(e).__name__}）")
            return job, e  # この1検索だけの失敗。続いていなければ、ほかの町名・検索語は続ける
        with lock:
            streak[0] = 0
            fetched_total[0] += len(got)
            if got:
                term_hit[term] = term_hit.get(term, 0) + 1
                term_zero[term] = 0
            else:
                term_zero[term] = term_zero.get(term, 0) + 1
            dead = all(term_zero.get(t, 0) >= ZERO_TERM_IN_A_ROW and not term_hit.get(t) for t in all_terms)
            tripped = dead and not stop.is_set()
            if tripped:
                why = f"すべての検索語が、町名を変えても{ZERO_TERM_IN_A_ROW}回続けて0件だったため、残りの検索は自動で止めました。検索語が地図に合っていないか、場所の書き方が合っていない可能性があります。取れた分でリストを作ります"
                auto_stop.append(why)
                RUN_STATS["auto_stopped_zero"] = why
                stop.set()
        finish(job, len(got))
        if tripped:
            log(f"\n  {auto_stop[0]}")
        return job, got

    workers = max(1, min(PARALLEL_MAX, workers))
    if jobs:
        minutes = estimate_minutes(len(jobs), limit, workers)
        log(search_plan_message(budget, len(jobs), minutes, workers))
        RUN_STATS["estimate_minutes"] = minutes
    with ThreadPoolExecutor(max_workers=workers) as ex:
        got = dict(ex.map(run, jobs))
    return {(ui, term, sub): (got[(term, sub)] if isinstance(got[(term, sub)], BaseException) else [dict(g) for g in got[(term, sub)]])
            for ui, term, sub in keys}


def main(budget=None):
    """探す処理の本体。途中の予期しない例外でも、この回の内訳の印を failed で書いてから上げる(前の回の印を残さない)"""
    RUN_ID["started"] = dt.datetime.now().isoformat(timespec="seconds")
    RUN_ID["id"] = f"{RUN_ID['started']}#{_os.getpid()}"
    RUN_STATS.clear()
    RUN_STATS["out_of_area"] = 0
    map_browser_source.RUN_COUNTS.clear()
    map_browser_source.RUN_WARNINGS.clear()
    try:
        return _main(budget)
    except BaseException as e:
        if not isinstance(e, SystemExit):
            write_run_summary("failed", error=f"{type(e).__name__}: {str(e)[:200]}")
        raise


def _main(budget=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--area")
    ap.add_argument("--industry", action="append")
    ap.add_argument("--keyword", action="append", help="Googleマップで探す言葉(何個でも)。Googleマップでは名前での絞り込みはしない(地図データ osm では名前に含むものだけ)")
    ap.add_argument("--limit", type=positive_int, default=200)
    ap.add_argument("--fetch-limit", type=positive_int, default=None, help="検索する数の上限(候補を集める全体の件数)。未指定なら既定値200で止める")
    ap.add_argument("--parallel", type=int, default=3, help="Googleマップを同時に何本取るか(1〜5。既定3。増やすとブロックされやすい)")
    ap.add_argument("--source", action="append", choices=["osm", "google", "google_places_text"], help="地図の取得元(何個でも。既定google。Yahoo!マップは外した)")
    ap.add_argument("--input", help="CSV/TXT of company names or URLs")
    ap.add_argument("--overture-file", help="Locally extracted Overture Places CSV/JSONL/GeoJSON")
    ap.add_argument("--overture-release", default="2026-09-23.1")
    ap.add_argument("--overture-license", default="CDLA-Permissive-2.0/Apache-2.0")
    ap.add_argument("--boundary-geojson", help="Optional ward boundary GeoJSON for strict point-in-polygon filtering")
    ap.add_argument("--portal", action="append", help="ポータルサイト等のURL(複数指定可)。掲載企業のHPを自動で集める")
    ap.add_argument("--include-large", action="store_true", help="互換のために残している(既定で含める。何も変わらない)")
    ap.add_argument("--exclude-large", action="store_true", help="大手・チェーンと推定した会社を外す(既定は含める)")
    ap.add_argument("--ng", help="送らない会社のドメイン一覧(1行1件のTXT/CSV)。既存客・競合を入れる")
    ap.add_argument("--ng-state-db", help="SQLite DB for NG and send-state gating")
    ap.add_argument("--resume-checks", action="store_true", help="サイト確認結果を既存SQLiteに保存し、途中停止後に行単位で再開する")
    ap.add_argument("--check-retry-limit", type=int, default=2, help="--resume-checks 時の取得再試行上限")
    ap.add_argument("--official-research", action="store_true", help="SNS/portal/no-contact URLs are re-searched for an official HP, then checked again")
    a = ap.parse_args()
    import license_check  # 画面を通らない直接実行でも、利用許可の確認を通す
    ok, why = license_check.gate()
    if not ok:
        print(why)
        return 7
    blocked = set()  # Googleに止められた取得元。ポータル指定の回でも、最後の内訳を書くために先に用意する
    clear_stop()  # 前の回に置かれたまま残った止める指示で、この回を止めない
    write_run_summary("running", area=a.area)  # 前の回の印を、この回の結果と取り違えさせない
    ng_store = open_ng_state_store(a.ng_state_db)

    if a.source and "google_places_text" in a.source:
        log(GOOGLE_PLACES_TEXT_BLOCK)
        return 1

    if a.portal:
        cap = max(a.fetch_limit, 1) if getattr(a, "fetch_limit", 0) else DEFAULT_FETCH_LIMIT  # ポータルから集める候補の上限も「検索する数の上限」にそろえる(止まる条件はこれ1つ)
        rows, labels, errors, seen_hosts = [], [], [], set()
        for purl in a.portal:
            if len(rows) >= cap:
                break
            try:
                got = portal_source.fetch(purl, cap - len(rows), log)
            except RuntimeError as e:
                errors.append(f"{purl}: {e}")
                log(f"ポータル {purl} は開けませんでした: {e}")
                continue
            for r in got:
                h = urlparse(r.get("url", "")).netloc.lower()
                if h and h not in seen_hosts:
                    seen_hosts.add(h)
                    rows.append(r)
            labels.append(urlparse(purl).netloc)
            log(f"ポータル {purl} から {len(got)} 件(累計 {len(rows)} 件)")
        if not rows:
            write_run_summary("failed", area=a.portal[0], fetched=0, error=(errors[0] if errors else "ポータルから1件も集まりませんでした")[:200])
            return 1
        label = "ポータル_" + "+".join(labels)
    elif a.input or a.overture_file:
        rows = read_input_file(a.input) if a.input else []
        label = Path(a.input).stem if a.input else "overture"
    else:
        if a.area and (a.industry or a.keyword):
            area, inds, limit = a.area, a.industry or [], a.limit
        else:
            area, inds, limit = ask_interactive()
        a.limit, label, rows = limit, area, []
        RUN_STATS["area_notes"] = subareas.notes(area)
        for note in RUN_STATS["area_notes"]:
            log(f"  [地名] {note}")
        renamed = [(i, industries.canonical(i)) for i in inds if industries.canonical(i) != i]
        for old, new in renamed:
            log(f"  [業種] 「{old}」は、いまの分類「{new}」にまとめました。「{new}」の検索語で探します")
        inds = list(dict.fromkeys(industries.canonical(i) for i in inds))
        for ind in inds:
            if ind not in INDUSTRIES and ind not in MAP_TERMS:
                log(f"  [業種] 「{ind}」は業種の一覧にありません。この言葉をそのままGoogleマップで探します（地図で引きにくい言葉だと、関係の薄い店が混ざるか、ほとんど出ません）")
        units = [(i, None) for i in inds] + [("キーワード", k) for k in (a.keyword or [])]
        until = google_block_active() if "google" in (a.source or ["google"]) else ""
        if until:
            blocked.add("google")
            log(f"\nGoogleマップは、同時に動いていた別の検索がGoogleに止められたため、{until[11:16]}ごろまで使いません。時間をおいて探し直してください")
        pre = prefetch_google(area, units, a.limit, a.parallel, block_check=google_block_active, stop_check=stop_requested, fetch_limit=a.fetch_limit) if "google" in (a.source or ["google"]) and not until else {}
        for ui, (ind, kw) in enumerate(units):
            log(f"\n{area} で " + (f"「{kw}」を探しています（Googleマップでは名前での絞り込みはしません。関係の薄い店も混ざります）..." if kw else f"{ind} を地図データから探しています..."))
            for src in a.source or ["google"]:
                if src in blocked:
                    continue
                if src != "osm" and not kw and ind in NO_MAP:
                    log(f"  [{src}] 「{ind}」は地図で引ける言葉がないため、飛ばします")
                    continue
                try:
                    if src == "osm":
                        got = osm_source.fetch_keyword(area, kw, log) if kw else osm_source.fetch(area, ind, log)
                        got = limit_osm_rows(got, a.limit, log)
                    else:
                        got = []
                        subs = subareas.split(area)
                        if len(subs) == 1:
                            log(f"  {area} は町名の分割表が無いため、区・市全体を1つの検索で探します。地図は1検索で数十件までしか出ないため、取りこぼしがあります")
                        failed, first_cause, total = [], None, 0
                        not_run = 0
                        for term in ([kw] if kw else map_terms(ind)):
                            for sub in subs:
                                part = pre[(ui, term, sub)]
                                total += 1
                                if isinstance(part, NotRun):  # 止める指示・0件が続いた自動停止で実行しなかった分。失敗には数えない
                                    not_run += 1
                                    continue
                                if isinstance(part, BaseException):
                                    failed.append((term, sub, part))
                                    if first_cause is None and not str(part).startswith("ほかの検索が止まった"):
                                        first_cause = part  # 連鎖で止まった分ではなく、最初に失敗した原因の原文を残す
                                    continue
                                for g in part:
                                    g["area"] = area  # 町名は検索のためだけに使う。エリアの欄は元の指定のまま
                                got += part
                        if not_run:
                            RUN_STATS["not_run"] = RUN_STATS.get("not_run", 0) + not_run
                            log(f"  [{src}] 止めたため実行しなかった検索: {not_run}件（失敗ではありません）")
                        if failed:
                            RUN_STATS["fetch_failed"] = RUN_STATS.get("fetch_failed", 0) + len(failed)
                            log(f"  [{src}] {total}件の検索のうち {len(failed)}件が失敗しました。取れた分は残します。失敗した検索: " + "、".join(f"{sub}（{type(e).__name__}）" for _t, sub, e in failed[:5]) + (" ほか" if len(failed) > 5 else ""))
                            if first_cause is not None:
                                log(f"  [{src}] 最初の失敗の原文: {first_cause}")
                            hit = next((e for _t, _s, e in failed if isinstance(e, map_browser_source.GoogleBlocked)), None)
                            if hit is not None:
                                if src not in blocked and not isinstance(hit, BlockedElsewhere):
                                    mark_google_blocked(hit)  # 同時に動いている別のサーバーにも知らせる
                                blocked.add(src)
                                log(f"  Googleに止められたため、この回ではGoogleマップをもう使いません（ほかのサーバーも{BLOCK_MINUTES}分は新しく探しません）。取れた分だけで続けます")
                            if not got:
                                raise first_cause or failed[0][2]
                        for g in got:
                            g["industry"] = kw or ind  # 業種の欄は元の業種名のまま(文面・メール型の選択に使う)
                except Exception as e:  # ValueError・RuntimeError・playwright の例外・未インストール等。他の取得元の結果は残す
                    log(f"  [{src}] 止まりました: {e}")
                    if src == "google":
                        blocked.add(src)  # ブロックされた可能性。この回はもう Google を叩かない
                        log("  Googleマップは、この回ではもう使いません")
                    continue
                kept = [g for g in got if not out_of_area(g, area)]
                if len(kept) != len(got):
                    outside = [g for g in got if out_of_area(g, area)]
                    RUN_STATS["out_of_area"] += len(outside)
                    log(f"  [{src}] 住所が「{area}」の外にある店 {len(outside)}件を除きました（例: " + "、".join(f"{g['name'][:12]}［{g.get('address', '')[:18]}］" for g in outside[:3]) + "）")
                    got = kept
                log(f"  [{src}] {len(got)}件見つかりました（うちホームページが分かる会社 {sum(1 for g in got if g['url'])}件）")
                rows += got
    if a.overture_file:
        got = overture_source.fetch_file(a.overture_file, area=a.area or "東京都江東区",
                                         release=a.overture_release, license_text=a.overture_license,
                                         boundary_geojson=a.boundary_geojson, log=log)
        rows += got
        label = (label + "_overture") if label != "overture" else label
    rows, dup_places = dedupe_places(rows)
    RUN_STATS["duplicate_places"] = dup_places
    if dup_places:
        log(f"\n同じ店が複数の町名・検索語に出た分 {dup_places}件を、1件にまとめました")
    m = map_counts()
    if m["searches"]:
        log(f"\n地図の検索 {m['searches']}回: 取れた店 {m['rows']}件 / 読めなかった店 {m['unreadable']}件・名前を読めなかった店 {m['bad_name']}件 / 上限の件数に届く前に一覧が止まった検索 {m['cut_short']}回 / "
            f"結果1件で施設ページに直行して読めなかった検索 {m['place_missed']}回 / 結果0件の検索 {m['empty']}回")
    fetched = len(rows)
    if fetched == 0:
        failed_runs = RUN_STATS.get("fetch_failed", 0) > 0 or bool(blocked)
        if stop_requested():
            outcome = "stopped"
            log("\n止める指示で止めました。")
        elif failed_runs:
            outcome = "failed"
            log("\n会社が取れませんでした。これは「該当なし」ではなく、取得の失敗（通信の失敗かGoogleに止められた）です。上の失敗の原文を確かめて、時間をおいて探し直してください。")
        elif RUN_STATS.get("out_of_area"):
            outcome = "all_outside"
            log(f"\n地図では {RUN_STATS['out_of_area']}件見つかりましたが、住所が「{a.area}」と合わず、全部エリア外として外しました。0件ではありません。"
                "地名の書き方を確かめてください（例: 東京都江東区、東京都墨田区）。")
        elif RUN_STATS.get("auto_stopped_zero"):
            outcome = "auto_stopped"
            log("\n0件が続いたため自動で止めました。すべての検索を最後まで実行したわけではないので、0件とは言い切れません。場所の書き方（市区町村まで）と、業種・キーワードを変えて探し直してください。")
        elif missed_total(m):
            outcome = "missed"
            log("\n残った会社は0件ですが、読めなかった店・施設ページに直行して読めなかった検索があります。0件とは言い切れません。時間をおいて探し直してください。")
        else:
            outcome = "zero"
            log("\n会社が1件も見つかりませんでした。エリア名を市区町村まで入れる、業種を変える、を試してください。")
        write_run_summary(outcome, area=a.area, fetched=0, blocked=sorted(blocked))
        return 1
    reject_google_places_text_rows(rows)
    assign_budget_line_ids(rows)
    missing = sum(1 for r in rows if not r["url"])
    if missing and stop_requested():
        log(f"\n止める指示があったため、ホームページが分からない会社 {missing}件の検索はしません")
    elif missing:
        log(f"\nホームページが地図に載っていない会社 {missing}件を、会社名で検索して探します（1件あたり十数秒。電話と住所でその会社のHPかを照合します）...")
        n, tried = hp_finder.fill(rows, log, want=max(a.limit * 2, 20), candidate_allowed=lambda u: candidate_allowed(u, ng_store), budget=budget)
        log(f"  {n}件のホームページが分かりました（今回あらたに検索 {tried}件）")
    size_guess.mark(rows)  # 町名をまたいだ同じ店は、上の dedupe_places で先に1件にしてある
    targets, no_url = dedupe(rows)
    no_url_first = len(no_url)  # 段階別の件数(画面の「内訳」)に使う。ここで数える
    deduped_before_limit = list(targets)
    dup = sum(1 for r in rows if r["url"]) - len(targets)
    if dup > 0:
        log(f"同じホームページで同じ店(同じ電話、または同じ名前と住所)として1行にまとめた会社: {dup}件")
    if a.ng:
        targets, ng_n = apply_ng(targets, a.ng)
        log(f"NGリストに該当して除外: {ng_n}件")
    before = len(targets)
    targets = [r for r in targets if not is_ng_url(r["url"], ng_store)]  # configured NG source
    if before != len(targets):
        log(f"自動NGリストに該当して除外: {before - len(targets)}件")
    big = [r for r in targets if r["size"] != "中小(推定)"]
    large_excluded = len(big) if a.exclude_large else 0
    if big and a.exclude_large:
        targets = [r for r in targets if r["size"] == "中小(推定)"]
        log(f"大手・チェーンと推定して外した会社: {len(big)}社（例: " + "、".join(f"{r['name']}［{r['size_why']}］" for r in big[:4]) + "）。含めたいときは「大手・チェーンも含める」を入れてください")
        if not targets:
            log(f"\n取れた {len(big)}社は、すべて大手・チェーンと推定して外しました。0件ではありません。「大手・チェーンも含める」を入れて、探し直してください。")
            write_run_summary("all_large", area=a.area, fetched=fetched, blocked=sorted(blocked))
            return 1
    else:
        targets.sort(key=lambda r: r["size"] != "中小(推定)")  # 含める場合も、中小を先に調べる
    cut = max(0, len(targets) - a.limit)
    targets = targets[:max(1, a.limit)]
    if cut:
        log(f"上限{a.limit}社を超えた {cut}社は調べていません（後ろの業種・取得元から落ちます）")
    log(f"\n{len(targets)}社のホームページで、問い合わせフォームを探します（1社あたり数秒〜十数秒）...")
    checked = run_checks(targets, resume_checks=a.resume_checks, retry_limit=a.check_retry_limit,
                         official_research=a.official_research, budget=budget, ng_store=ng_store, stop_check=stop_requested)
    if RUN_STATS.get("not_checked_stopped"):
        log(f"止める指示で止めました。まだ調べていなかった {RUN_STATS['not_checked_stopped']}社はリストに入れていません")
    checked, rededuped = dedupe_checked_by_url(checked)
    if rededuped:
        no_url += rededuped
        log(f"official HP re-search duplicate removed: {len(rededuped)}")
    final_ng = [r for r in checked if is_output_ng(r, ng_store)]
    if final_ng:
        checked = [r for r in checked if not is_output_ng(r, ng_store)]
        no_url += final_ng
        log(f"final NG re-check removed: {len(final_ng)}")
    dropped = [r for r in checked if r.get("discard")]
    if dropped:
        checked = [r for r in checked if not r.get("discard")]
        no_url += dropped  # 「URLが分からない会社」の表に残す(検索用リンクで手で探せる)
        log(f"検索で見つけたHPが別会社の疑いで捨てた会社: {len(dropped)}社（例: " + "、".join(f"{r['name']}［{r['discard'][:30]}］" for r in dropped[:3]) + "）")
    refused = [r for r in checked if r["refuse"]]
    added = sum(add_auto_ng(r["url"], "refused_form: " + r["refuse_text"][:30], r["name"], ng_store) for r in refused)
    checked = [r for r in checked if not r["refuse"]]  # お断りの会社は一覧に出さず、NGリストへ入れる
    if refused:
        log(f"営業お断りの表示があった会社 {len(refused)}件を一覧から外しました（NGリストに新しく入れたのは {added}件）")
    summarize_source_funnel(rows, deduped_before_limit, checked, no_url)
    summarize_contact_states(checked)
    summarize(fetched, checked, no_url, refused=len(refused))
    x, c = write_outputs(checked, no_url, label)
    log(f"\n保存しました:\n  {x}\n  {c}")
    try:
        summarize_sendable(c, fetched, checked, no_url)
    except Exception as e:  # 件数の内訳が出せなくても、収集の結果は保存済み。理由を残す
        log(f"件数の内訳を出せませんでした: {e!r}")
    failed_some = RUN_STATS.get("fetch_failed", 0) > 0 or bool(blocked)
    outcome = "stopped" if RUN_STATS.get("not_checked_stopped") or stop_requested() else ("partial" if failed_some else "ok")
    write_run_summary(outcome, area=a.area, fetched=fetched, checked=len(checked), no_url=len(no_url),
                      dropped_other_company=len(dropped), refused=len(refused), blocked=sorted(blocked), list_file=Path(c).name,
                      not_checked_stopped=RUN_STATS.get("not_checked_stopped", 0),
                      funnel=build_funnel(fetched, no_url_first, len(deduped_before_limit), large_excluded, cut, len(targets), len(dropped), len(refused), checked))
    return 0


if __name__ == "__main__":
    sys.exit(main())



